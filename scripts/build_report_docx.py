"""產生正式專題報告書（Word）。

內容來源：docs/01_正式規格/正式報告書/報告書內容.md
圖檔來源：docs/images/（對應寫在下面的 FIGURES）
輸出：docs/01_正式規格/正式報告書/資管科智慧學習歷程系統_專題報告書.docx

格式依學長姐專題成果報告書（115 招生平台）實測：中文標楷體、英文 Times New Roman；
章標題 26 級、節標題 20 級、小標 18 級、內文 16 級、表格與圖說 14 級；每章換頁。
封面、目錄、圖目錄與頁尾頁碼由本程式產生。目錄是 Word 功能變數，
用 Word 開啟時會自動更新頁碼；上傳到 Google 文件後也會轉成文件目錄。

執行：在 D:\\democase\\ai 執行  .venv\\Scripts\\python.exe scripts\\build_report_docx.py
"""
import re
import struct
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "01_正式規格" / "正式報告書" / "報告書內容.md"
OUT = ROOT / "docs" / "01_正式規格" / "正式報告書" / "資管科智慧學習歷程系統_專題報告書.docx"
LITE_SRC = ROOT / "docs" / "01_正式規格" / "正式報告書" / "報告書內容_精簡版.md"
LITE_OUT = ROOT / "docs" / "01_正式規格" / "正式報告書" / "資管科智慧學習歷程系統_專題報告書_精簡版.docx"
ANNEX_SRC = ROOT / "docs" / "01_正式規格" / "正式報告書" / "開發規格附件內容.md"
ANNEX_OUT = ROOT / "docs" / "01_正式規格" / "正式報告書" / "資管科智慧學習歷程系統_系統開發規格附件.docx"
IMG = ROOT / "docs" / "images"

CJK_FONT, LATIN_FONT = "標楷體", "Times New Roman"
SIZE_BODY, SIZE_TABLE, SIZE_H1, SIZE_H2, SIZE_H3 = 16, 14, 26, 20, 18

COVER = {
    "school": "康寧學校財團法人康寧大學",
    "dept": "資訊管理科",
    "kind": "畢業專題",
    "title": "資管科智慧學習歷程系統",
    "students": [("111534105", "田心于"), ("111534107", "張馨羽"), ("111534124", "林承諺")],
    "advisor": "李栢松 老師",
    "date": "中 華 民 國  116  年 06 月",
}

# 圖號 → 圖檔（與 docs/04_圖表與介面/03_檢查與圖表規範/diagrams.md、流程圖大全簡報一致）
FIGURES = {
    "1": "D-01-use-case-student.png", "2": "D-01-use-case-teacher.png", "3": "D-01-use-case-staff.png",
    "4": "D-01-use-case-admin.png", "5": "D-02-login-role-routing.png", "6": "D-10-main-line-flow.png",
    "7": "D-21-school-data-import-flow.png", "8": "D-16-teacher-confirm-flow.png",
    "9": "D-18-staff-service-record-flow.png", "10": "D-20-form-download-flow.png",
    "11": "D-13-admission-rule-update-flow.png", "12": "D-19-admission-announcement-flow.png",
    "13": "D-22-info-source-classify-flow.png", "14": "D-17-admin-maintain-requirements.png",
    "15": "D-14-role-responsibility.png", "16": "D-05-system-architecture.png", "17": "D-06-dfd.png",
    "18-1": "D-09-erd-part1-learning.png", "18-2": "D-09-erd-part2-review.png",
    "18-3": "D-09-erd-part3-sources.png", "19": "D-15-report-overview.png",
    "20": "D-23-function-architecture.png", "21": "D-24-gantt-2026.png", "22": "D-25-login-flow.png",
    "23": "D-26-student-upload-flow.png", "24": "D-27-target-recommendation-flow.png",
    "25": "D-28-review-package-flow.png", "26": "D-29-certificate-flow.png",
    "27": "D-30-admin-source-flow.png", "28": "D-31-student-role-architecture.png",
    "29": "D-32-teacher-role-architecture.png", "30": "D-33-staff-role-architecture.png",
    "31": "D-34-admin-role-architecture.png", "32": "D-35-permission-matrix.png",
    "33": "D-36-student-detailed-flow.png", "34": "D-37-teacher-detailed-flow.png",
    "35": "D-38-staff-detailed-flow.png", "36": "D-39-admin-detailed-flow.png",
    "37": "D-40-version-history-download-flow.png", "38": "D-41-student-dashboard-wireframe.png",
    "39": "D-42-data-confirmation-wireframe.png", "40": "D-43-teacher-review-wireframe.png",
    "41": "D-44-admin-maintenance-wireframe.png",
}

# A4 直式可用寬高（公分）；太寬的圖改放橫式頁面
PORTRAIT_W, PORTRAIT_H = 15.9, 21.5
LANDSCAPE_W, LANDSCAPE_H = 24.6, 13.5
WIDE_RATIO = 2.4   # 寬／高超過這個比例就放橫式頁


def set_fonts(run_or_style, size=None, bold=None):
    font = run_or_style.font
    font.name = LATIN_FONT
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    rpr = run_or_style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for key in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(key), LATIN_FONT)
    fonts.set(qn("w:eastAsia"), CJK_FONT)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    set_fonts(normal, SIZE_BODY)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Heading 1", SIZE_H1), ("Heading 2", SIZE_H2), ("Heading 3", SIZE_H3)):
        st = doc.styles[name]
        set_fonts(st, size, True)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.italic = False
        pf = st.paragraph_format
        pf.keep_with_next = True
        pf.space_before = Pt(12)
        pf.space_after = Pt(8)
        if name == "Heading 1":
            pf.page_break_before = True


def add_runs(par, text, size=None, bold=None):
    """把 **粗體** 標記轉成 run；<br> 轉成換行。"""
    for i, part in enumerate(re.split(r"\*\*", text)):
        for j, piece in enumerate(part.split("<br>")):
            if j:
                par.add_run().add_break()
            if piece:
                run = par.add_run(piece)
                set_fonts(run, size, True if (i % 2 == 1) else bold)


def add_field(par, code, placeholder=""):
    run = par.add_run()
    for tag, attr in (("w:fldChar", "begin"),):
        el = OxmlElement(tag); el.set(qn("w:fldCharType"), attr); run._r.append(el)
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = code
    run._r.append(instr)
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate"); run._r.append(sep)
    if placeholder:
        t = OxmlElement("w:t"); t.text = placeholder; run._r.append(t)
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end"); run._r.append(end)
    set_fonts(run)
    return run


REF = re.compile(r"圖 (\d+(?:-\d+)?)(?:～(\d+(?:-\d+)?))?")


def renumber_refs(text, display):
    """把內文的「圖 N」「圖 N～M」換成報告書中依出現順序的圖號。"""
    order = list(FIGURES)

    def repl(m):
        a, b = m.group(1), m.group(2)
        if a not in display or (b and b not in display):
            return m.group(0)
        ids = order[order.index(a): order.index(b) + 1] if b else [a]
        nums = sorted(display[i] for i in ids if i in display)
        if len(nums) > 1 and nums == list(range(nums[0], nums[-1] + 1)):
            return f"圖 {nums[0]}～{nums[-1]}"
        return "、".join(f"圖 {n}" for n in nums)
    return REF.sub(repl, text)


def png_size(path):
    with open(path, "rb") as f:
        head = f.read(24)
    w, h = struct.unpack(">II", head[16:24])
    return w, h


def set_orientation(section, landscape):
    w, h = section.page_width, section.page_height
    if landscape != (w > h):
        section.page_width, section.page_height = h, w
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT


def add_figure(doc, num, caption, state):
    path = IMG / FIGURES[num]
    w, h = png_size(path)
    wide = w / h >= WIDE_RATIO
    if wide != state["landscape"]:
        sec = doc.add_section(WD_SECTION.NEW_PAGE)
        set_orientation(sec, wide)
        state["landscape"] = wide
    max_w, max_h = (LANDSCAPE_W, LANDSCAPE_H) if wide else (PORTRAIT_W, PORTRAIT_H)
    scale = min(max_w / w, max_h / h)
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(str(path), width=Cm(w * scale), height=Cm(h * scale))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(cap, f"圖 {state['display'][num]}：{caption}", SIZE_TABLE)


TALL_RATIO = 0.75  # 寬／高小於這個比例的窄長圖，連續兩張時並排


def is_tall(num):
    w, h = png_size(IMG / FIGURES[num])
    return w / h < TALL_RATIO


def add_figure_pair(doc, a, b, state):
    """兩張窄長圖並排放在同一頁（無框線表格，各自附圖說）。"""
    if state["landscape"]:
        back_to_portrait(doc, state)
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, (num, caption) in enumerate((a, b)):
        path = IMG / FIGURES[num]
        w, h = png_size(path)
        scale = min(7.6 / w, PORTRAIT_H / h)
        p = table.cell(0, c).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(w * scale), height=Cm(h * scale))
        q = table.cell(1, c).paragraphs[0]
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(q, f"圖 {state['display'][num]}：{caption}", SIZE_TABLE)
    doc.add_paragraph()


def back_to_portrait(doc, state):
    if state["landscape"]:
        sec = doc.add_section(WD_SECTION.NEW_PAGE)
        set_orientation(sec, False)
        state["landscape"] = False


def shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            add_runs(par, text, SIZE_TABLE, True if r == 0 else None)
            if r == 0:
                shade(cell, "D9D9D9")
    # 欄寬依各欄最長的一行文字分配（短欄給窄、長欄給寬），總寬等於版心寬
    longest = [max(max(len(x) for x in re.sub(r"\*\*", "", row[c]).split("<br>")) for row in rows)
               for c in range(len(rows[0]))]
    weights = [min(max(n, 4), 30) for n in longest]
    table.autofit = False
    for c, wgt in enumerate(weights):
        width = Cm(PORTRAIT_W * wgt / sum(weights))
        for row in table.rows:
            row.cells[c].width = width
    trpr = table.rows[0]._tr.get_or_add_trPr()   # 表頭跨頁重複
    hdr = OxmlElement("w:tblHeader"); hdr.set(qn("w:val"), "true"); trpr.append(hdr)
    doc.add_paragraph()


def cover(doc):
    def line(text, size, bold=False, before=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        add_runs(p, text, size, bold)
    line(COVER["school"], 24, True, 60)
    line(COVER["dept"], 24, True, 12)
    line(COVER["kind"], 24, True, 12)
    line(COVER["title"], 36, True, 90)
    table = doc.add_table(rows=len(COVER["students"]) + 1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (sid, name) in enumerate(COVER["students"]):
        for c, text in enumerate(("專 題 生：" if i == 0 else "", sid, name)):
            p = table.cell(i, c).paragraphs[0]; add_runs(p, text, 18)
    last = len(COVER["students"])
    add_runs(table.cell(last, 0).paragraphs[0], "指導老師：", 18)
    add_runs(table.cell(last, 1).paragraphs[0], COVER["advisor"], 18)
    for row in table.rows:
        row.cells[0].width = Cm(4); row.cells[1].width = Cm(4); row.cells[2].width = Cm(4)
    doc.paragraphs[0].paragraph_format.space_before = Pt(60)
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(120)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(p, COVER["date"], 20, True)


def toc(doc, headings_only=True):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before = True
    add_runs(p, "目錄", SIZE_H1, True)
    add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u', "（開啟 Word 後會自動更新目錄與頁碼）")


def footer_page_numbers(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True     # 封面不放頁碼
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(p, "第 ", 12); add_field(p, "PAGE", "1").font.size = Pt(12); add_runs(p, " 頁", 12)


def parse(md):
    md = re.sub(r"<!--.*?-->", "", md, flags=re.S)
    lines = md.splitlines()
    blocks, i = [], 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1; continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows)); continue
        m = re.match(r"^(#{1,3}) (.*)$", line)
        if m:
            blocks.append((f"h{len(m.group(1))}", m.group(2).strip()))
        elif re.match(r"^\[\[圖 (\S+?)：(.*)\]\]$", line):
            m = re.match(r"^\[\[圖 (\S+?)：(.*)\]\]$", line)
            blocks.append(("figure", (m.group(1), m.group(2))))
        elif line.startswith("- "):
            blocks.append(("bullet", line[2:].strip()))
        elif re.match(r"^\d+\. ", line):
            blocks.append(("numbered", line.strip()))
        else:
            blocks.append(("para", line.strip()))
        i += 1
    return blocks


def build(src=SRC, out=OUT, title=None):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.54))
    setup_styles(doc)
    cover(doc) if title is None else annex_cover(doc, title)
    toc(doc)
    blocks = parse(src.read_text(encoding="utf-8"))
    figures = [b[1] for b in blocks if b[0] == "figure"]
    missing = [n for n, _ in figures if n not in FIGURES or not (IMG / FIGURES[n]).exists()]
    if missing:
        raise SystemExit(f"找不到圖檔：圖 {missing}")
    display = {num: i for i, (num, _) in enumerate(figures, 1)}
    blocks = [(k, renumber_refs(v, display)) if isinstance(v, str) else
              (k, [[renumber_refs(c, display) for c in row] for row in v]) if k == "table" else (k, v)
              for k, v in blocks]
    if figures:   # 圖目錄（靜態清單，依新圖號）
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.page_break_before = True
        add_runs(p, "圖目錄", SIZE_H1, True)
        for num, cap in figures:
            q = doc.add_paragraph(); q.paragraph_format.space_after = Pt(0)
            add_runs(q, f"圖 {display[num]}　{cap}", SIZE_TABLE)
    state = {"landscape": False, "display": display}
    pair_start, pair_second = set(), set()   # 連續兩張窄長圖並排
    k = 0
    while k < len(blocks) - 1:
        (k1, v1), (k2, v2) = blocks[k], blocks[k + 1]
        if k1 == k2 == "figure" and is_tall(v1[0]) and is_tall(v2[0]):
            pair_start.add(k); pair_second.add(k + 1); k += 2
        else:
            k += 1
    for idx, (kind, val) in enumerate(blocks):
        if idx in pair_second:
            continue
        if idx in pair_start:
            add_figure_pair(doc, val, blocks[idx + 1][1], state)
            continue
        if kind != "figure":
            back_to_portrait(doc, state)
        if kind in ("h1", "h2", "h3"):
            h = doc.add_heading(level=int(kind[1]))
            add_runs(h, val)
        elif kind == "para":
            add_runs(doc.add_paragraph(), val)
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, val)
        elif kind == "numbered":
            p = doc.add_paragraph()
            num, rest = val.split(" ", 1)
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.9)
            add_runs(p, f"{num} {rest}")
        elif kind == "table":
            add_table(doc, val)
        elif kind == "figure":
            add_figure(doc, val[0], val[1], state)
    back_to_portrait(doc, state)
    # 每個新節沿用第一節的頁碼頁尾
    footer_page_numbers(doc)
    for s in doc.sections[1:]:
        s.footer.is_linked_to_previous = True
        s.different_first_page_header_footer = False
    # 開檔時自動更新目錄
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true"); settings.append(upd)
    # 清單樣式字型
    set_fonts(doc.styles["List Bullet"], SIZE_BODY)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"已產生 {out}（圖 {len(figures)} 張、分節 {len(doc.sections)}）")


def annex_cover(doc, title):
    for text, size, before in ((COVER["title"], 30, 160), (title, 26, 30), (COVER["date"], 20, 200)):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        add_runs(p, text, size, True)


if __name__ == "__main__":
    import sys
    if "--annex" in sys.argv:
        build(ANNEX_SRC, ANNEX_OUT, "系統開發規格附件")
    elif "--lite" in sys.argv:      # 精簡版：只放符合專題資格的核心圖
        build(LITE_SRC, LITE_OUT)
    else:
        build()

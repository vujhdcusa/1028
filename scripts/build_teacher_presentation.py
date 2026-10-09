from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / "docs" / "images"
OUT = ROOT / "docs" / "08_簡報" / "AI_指導老師討論簡報_2026-10-07.pptx"

NAVY = RGBColor(20, 46, 78)
BLUE = RGBColor(45, 115, 171)
TEAL = RGBColor(32, 137, 133)
ORANGE = RGBColor(221, 129, 47)
INK = RGBColor(35, 44, 54)
MUTED = RGBColor(93, 107, 121)
PALE = RGBColor(242, 246, 249)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(177, 63, 63)


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False,
             align=PP_ALIGN.LEFT, font="Microsoft JhengHei", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_lines(slide, lines, x, y, w, h, size=18, color=INK, leading=1.1):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.06)
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.space_after = Pt(8)
        p.line_spacing = leading
        p.font.name = "Microsoft JhengHei"
        p.font.size = Pt(size)
        p.font.color.rgb = color
    return box


def base_slide(prs, title, section=None, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY if dark else WHITE
    if not dark:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.18))
        bar.fill.solid(); bar.fill.fore_color.rgb = TEAL; bar.line.fill.background()
        add_text(slide, title, 0.62, 0.48, 11.9, 0.55, 27, NAVY, True)
        if section:
            add_text(slide, section.upper(), 11.35, 0.56, 1.35, 0.25, 9, TEAL, True, PP_ALIGN.RIGHT)
        add_text(slide, "資管科智慧學習歷程系統｜指導老師討論版", 0.62, 7.18, 8.5, 0.2, 8, MUTED)
    return slide


def add_image_contain(slide, path, x, y, w, h):
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    ratio = min(w / iw, h / ih)
    nw, nh = iw * ratio, ih * ratio
    left = x + (w - nw) / 2
    top = y + (h - nh) / 2
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(nw), height=Inches(nh))


def add_callout(slide, label, body, x, y, w, h, color=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = PALE
    shape.line.color.rgb = color; shape.line.width = Pt(1.2)
    add_text(slide, label, x + 0.18, y + 0.12, w - 0.35, 0.28, 13, color, True)
    add_text(slide, body, x + 0.18, y + 0.48, w - 0.35, h - 0.58, 16, INK)


def diagram_slide(prs, filename, title, note=None):
    slide = base_slide(prs, title, "流程圖附錄")
    path = IMG / filename
    if path.exists():
        add_image_contain(slide, path, 0.55, 1.25, 12.25, 5.65)
    if note:
        add_text(slide, note, 0.72, 6.86, 11.9, 0.25, 10, MUTED, False, PP_ALIGN.CENTER)
    return slide


def make_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = base_slide(prs, "資管科智慧學習歷程系統", dark=True)
    add_text(slide, "指導老師討論簡報", 0.82, 1.55, 7.4, 0.7, 34, WHITE, True)
    add_text(slide, "目前系統流程、AI Agent 與待確認事項", 0.86, 2.35, 8.0, 0.45, 22, RGBColor(196, 222, 229))
    add_text(slide, "報告目的：確認主線是否正確，再決定哪些功能進入下一階段。", 0.86, 4.95, 8.6, 0.45, 18, WHITE)
    add_text(slide, "2026 / 10 / 07", 0.86, 6.65, 2.5, 0.25, 12, RGBColor(196, 222, 229))
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.55), Inches(0.8), Inches(1.1), Inches(5.9))
    accent.fill.solid(); accent.fill.fore_color.rgb = TEAL; accent.line.fill.background()
    add_text(slide, "學習\n歷程", 10.74, 2.6, 0.75, 1.4, 22, WHITE, True, PP_ALIGN.CENTER)

    slide = base_slide(prs, "今天要請老師確認的三件事", "報告重點")
    add_callout(slide, "01｜主線", "學生如何從資料整理，走到一份可送老師確認的備審資料。", 0.75, 1.55, 3.75, 2.2, BLUE)
    add_callout(slide, "02｜權限與 AI", "學生、導師、管理員各自能看什麼；AI 可以使用到什麼程度。", 4.78, 1.55, 3.75, 2.2, TEAL)
    add_callout(slide, "03｜未決策", "哪些規則需要科辦或老師確認，避免我們把錯誤假設寫進系統。", 8.81, 1.55, 3.75, 2.2, ORANGE)
    add_text(slide, "簡報原則", 0.78, 4.45, 1.8, 0.35, 16, NAVY, True)
    add_rich_lines(slide, ["先確認主線，再擴充功能", "AI 只提供草稿與查詢，不取代學生與老師判斷", "所有招生規則先由管理者確認後才發布"], 0.78, 4.9, 11.5, 1.4, 20)

    slide = base_slide(prs, "系統要解決的核心問題", "系統定位")
    add_text(slide, "學生資料分散，老師難以逐一檢查，招生要求每年也會變動。", 0.82, 1.35, 11.6, 0.5, 23, NAVY, True)
    add_callout(slide, "學生端", "整理五年學習歷程\n選擇目標校系\n看見缺件與下一步", 0.82, 2.25, 3.65, 2.45, BLUE)
    add_callout(slide, "教師端", "看本班學生資料\n查看原始檔\n確認或要求修改備審內容", 4.84, 2.25, 3.65, 2.45, TEAL)
    add_callout(slide, "管理端", "維護校系招生規則\n審核 Agent 擷取結果\n發布可用版本", 8.86, 2.25, 3.65, 2.45, ORANGE)
    add_text(slide, "系統輸出：學生可下載一份依目標校系整理的備審資料 Word。", 0.82, 5.45, 11.5, 0.45, 21, INK, True, PP_ALIGN.CENTER)

    slide = base_slide(prs, "目前確認的主線流程", "核心流程")
    add_image_contain(slide, IMG / "D-10-main-line-flow.png", 0.45, 1.15, 12.45, 5.8)
    add_text(slide, "學生選校系 → 上傳六類資料 → 缺件比對 → 自傳草稿 → 導師確認 → 下載備審資料", 0.75, 6.92, 11.8, 0.28, 12, NAVY, True, PP_ALIGN.CENTER)

    slide = base_slide(prs, "角色分工：每個人只做自己的事", "角色與權限")
    add_image_contain(slide, IMG / "D-14-role-responsibility.png", 0.65, 1.2, 5.0, 5.65)
    add_rich_lines(slide, ["學生：建立與修正自己的學習歷程", "導師：查看本班，確認內容是否合理", "行政／管理員：維護規則與招生資料", "AI Agent：擷取、整理、提醒，不直接發布"], 6.05, 1.55, 6.3, 3.25, 20)
    add_callout(slide, "要請老師確認", "導師是否需要看到成績單、證照與其他原始檔？專任教師與科助是否需要不同權限？", 6.05, 5.1, 6.2, 1.15, ORANGE)

    for filename, title, role in [
        ("D-01-use-case-student.png", "學生端使用案例", "學生"),
        ("D-01-use-case-teacher.png", "導師端使用案例", "導師"),
        ("D-01-use-case-staff.png", "科辦／科助使用案例", "科辦／科助"),
        ("D-01-use-case-admin.png", "管理端使用案例", "行政／管理員"),
    ]:
        slide = base_slide(prs, title, f"{role}端")
        add_image_contain(slide, IMG / filename, 0.55, 1.12, 12.2, 5.9)

    slide = base_slide(prs, "競賽報名流程（後續擴充草稿）", "待問科助")
    add_text(slide, "以下是討論草稿，不列入目前主線；是否由科辦開放報名，仍要先問科助。", 0.82, 1.22, 11.7, 0.42, 20, ORANGE, True, PP_ALIGN.CENTER)
    competition_steps = [
        ("科辦", "建立競賽事項\n報名期限、主辦方、報名方式"),
        ("學生", "從系統申請報名\n留下參賽紀錄"),
        ("指導老師", "確認參賽資格或\n指導老師關係"),
        ("主辦方／AREC", "提供參賽或得獎\n證明"),
        ("科辦", "管理公假、嘉獎\n等校內行政紀錄"),
        ("學生", "上傳證明\n納入學習歷程"),
    ]
    for i, (actor, action) in enumerate(competition_steps):
        x = 0.55 + (i % 3) * 4.25
        y = 1.95 + (i // 3) * 1.75
        color = [ORANGE, BLUE, TEAL, NAVY, ORANGE, BLUE][i]
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(3.55), Inches(1.18))
        shape.fill.solid(); shape.fill.fore_color.rgb = PALE
        shape.line.color.rgb = color; shape.line.width = Pt(1.4)
        add_text(slide, actor, x + 0.15, y + 0.12, 3.25, 0.27, 14, color, True, PP_ALIGN.CENTER)
        add_text(slide, action, x + 0.18, y + 0.47, 3.18, 0.55, 15, INK, False, PP_ALIGN.CENTER)
    add_callout(slide, "草稿中的來源邊界", "外部競賽證明由主辦單位或 AREC 開立；科辦只有在自己是主辦方時才可能開立證明。公假與嘉獎是另一類校內行政紀錄。以上仍待科助確認。", 0.8, 5.65, 11.7, 0.88, ORANGE)

    slide = base_slide(prs, "AI Agent：系統的兩個 AI 角色", "AI 功能")
    add_callout(slide, "資管科小助手", "查詢科上規則\n查看自己的進度與缺件\n回答下一步問題\n檢查自傳內容一致性", 0.9, 1.55, 5.45, 2.55, TEAL)
    add_callout(slide, "自傳小幫手", "一次問四題\n依學生回答產生草稿\n空白內容標示待補\n學生修改後再送導師確認", 6.95, 1.55, 5.45, 2.55, ORANGE)
    add_callout(slide, "共同邊界", "AI 提供草稿、分類與提醒；學生確認資料，導師確認備審，管理者確認招生規則。", 0.9, 4.65, 11.5, 1.15, NAVY)

    slide = base_slide(prs, "招生資料 Agent：每年更新，但不自動發布", "AI 功能")
    add_image_contain(slide, IMG / "D-13-admission-rule-update-flow.png", 0.45, 1.15, 7.0, 5.75)
    add_rich_lines(slide, ["1. 管理者提供官方來源", "2. Agent 偵測變更並擷取校系要求", "3. 系統標記為 candidate", "4. 管理者核對後才發布 published", "5. 學生只看到已發布版本"], 7.7, 1.55, 4.9, 3.5, 18)
    add_callout(slide, "要請老師確認", "哪些官方網址可以作為來源？每年由誰核對與發布？", 7.7, 5.35, 4.9, 1.0, RED)

    slide = base_slide(prs, "資料流與系統架構", "技術概要")
    add_image_contain(slide, IMG / "D-05-system-architecture.png", 0.35, 1.15, 6.15, 5.8)
    add_image_contain(slide, IMG / "D-06-dfd.png", 6.75, 1.15, 6.25, 5.8)

    slide = base_slide(prs, "資料表目前如何支撐主線", "資料設計")
    add_image_contain(slide, IMG / "D-09-erd-v2.png", 0.45, 1.05, 12.35, 5.95)
    add_text(slide, "主線資料：學生、原始檔、學習紀錄、目標校系、備審輸出、導師審核", 0.7, 6.9, 11.9, 0.28, 12, MUTED, False, PP_ALIGN.CENTER)

    slide = base_slide(prs, "目前 AI 使用的安全邊界", "風險與限制")
    add_rich_lines(slide, ["學生自己的 API Key：只在這次請求使用，不存檔、不寫入資料庫、不放進網址", "系統知識庫：先由科上文件整理，AI 回答時附來源，仍需要老師或科辦核對", "AI 產出：是草稿與提醒，學生要修改，導師要確認，系統才產出正式 Word", "招生規則：Agent 只能產生待審版本，不能直接覆蓋已發布規則"], 0.85, 1.45, 11.7, 3.9, 21)
    add_callout(slide, "需要決策", "是否允許把學生資料送到 Google Gemini？若允許，採免費方案、付費方案，或要求學生同意？", 0.85, 5.55, 11.55, 0.9, RED)

    slide = base_slide(prs, "請老師先確認：會擋住開發的問題", "討論清單 1")
    add_text(slide, "這三題若沒有答案，AI 與權限設計都可能重做。", 0.85, 1.22, 11.7, 0.4, 22, RED, True)
    add_rich_lines(slide, ["1. 學生資料可以送給 Google Gemini 嗎？需要同意書或改用付費方案嗎？", "2. AI 費用由學校、科上、專題經費，還是由學生自行提供金鑰？", "3. AI 協助自傳是否可放入備審？是否要標示「AI 協助草稿」並由導師確認？"], 0.95, 2.0, 11.4, 2.7, 23)
    add_text(slide, "建議請老師直接給「可以／不可以／需要補件」的決定。", 0.95, 5.75, 11.3, 0.45, 19, NAVY, True, PP_ALIGN.CENTER)

    slide = base_slide(prs, "請老師或科辦確認：資料與規則", "討論清單 2")
    add_rich_lines(slide, ["4. 9 份知識庫文件由誰核對？哪些學年度、哪些屆別適用？", "5. 校系的「個人特質與專業技能」要對應哪些資料類別？例如 50% 權重怎麼算？", "6. 招生 Agent 要監測哪些官方網址？各校分則是否列入？", "7. 導師是否可看本班學生的成績單、證照、佐證照片與自傳原始檔？專任教師是否不同？"], 0.95, 1.35, 11.4, 3.8, 20)
    add_callout(slide, "目前已知", "系統可以做到權限限制，但「誰應該看得到」必須由老師與科辦決定。", 0.95, 5.55, 11.4, 0.9, TEAL)

    slide = base_slide(prs, "請老師補充：尚未定義的角色與例外", "討論清單 3")
    add_rich_lines(slide, ["8. 科助在系統裡的任務是什麼？要維護資料、協助審核，還是只查看統計？", "9. 期中修完的課暫時看不到、微學分不在成績單上，要怎麼補登與驗證？", "10. 讀書計畫是否也需要 AI 小幫手？還是先完成自傳即可？", "11. 下次驗收要看完整示範流程，還是先看主線功能？"], 0.95, 1.45, 11.4, 3.8, 20)
    add_text(slide, "這些可以晚一點做，但要先留下明確的待辦與負責人。", 0.95, 5.75, 11.4, 0.45, 19, ORANGE, True, PP_ALIGN.CENTER)

    slide = base_slide(prs, "建議現場示範順序", "展示腳本")
    steps = ["學生登入", "選目標校系", "上傳六類資料", "查看缺件", "自傳小幫手", "小助手問答", "導師確認", "下載 Word"]
    for i, step in enumerate(steps):
        x = 0.75 + (i % 4) * 3.1
        y = 1.55 + (i // 4) * 2.0
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.45), Inches(0.95))
        shape.fill.solid(); shape.fill.fore_color.rgb = PALE
        shape.line.color.rgb = [BLUE, TEAL, ORANGE, NAVY][i % 4]
        add_text(slide, f"{i+1:02d}", x + 0.12, y + 0.17, 0.45, 0.35, 16, TEAL, True)
        add_text(slide, step, x + 0.58, y + 0.2, 1.65, 0.35, 18, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, "管理員補充示範：上傳招生資料 → Agent 擷取 → 管理者核對 → 發布校系規則", 0.9, 6.05, 11.4, 0.38, 18, INK, True, PP_ALIGN.CENTER)

    slide = base_slide(prs, "希望老師今天帶走的結論", "收尾")
    add_rich_lines(slide, ["主線：學生資料整理與備審產出流程是否正確？", "權限：導師、專任教師、科助與管理員的可見範圍如何分？", "AI：資料是否可送出、誰負責核對、AI 草稿是否可接受？", "下一步：哪些決策今天定案，哪些列為後續功能？"], 1.0, 1.55, 11.2, 3.4, 24, NAVY)
    add_text(slide, "謝謝老師，請協助我們先把主線定下來。", 0.9, 5.65, 11.5, 0.55, 25, TEAL, True, PP_ALIGN.CENTER)

    appendix = [
        ("D-01-use-case-v2.png", "附錄｜整體使用案例圖"),
        ("D-01-use-case-student.png", "附錄｜學生使用案例圖"),
        ("D-01-use-case-teacher.png", "附錄｜導師使用案例圖"),
        ("D-01-use-case-staff.png", "附錄｜科辦／科助使用案例圖"),
        ("D-01-use-case-admin.png", "附錄｜管理員使用案例圖"),
        ("D-02-login-role-routing.png", "附錄｜登入與角色導向流程"),
        ("D-05-system-architecture.png", "附錄｜系統架構圖"),
        ("D-06-dfd.png", "附錄｜資料流程圖"),
        ("D-09-erd-v2.png", "附錄｜ERD 資料表關聯"),
        ("D-10-main-line-flow.png", "附錄｜主線流程"),
        ("D-13-admission-rule-update-flow.png", "附錄｜招生規則更新流程"),
        ("D-14-role-responsibility.png", "附錄｜角色責任圖"),
        ("D-15-report-overview.png", "附錄｜報表總覽（第二版）"),
        ("D-16-teacher-confirm-flow.png", "附錄｜導師確認流程"),
        ("D-17-admin-maintain-requirements.png", "附錄｜管理員維護招生要求流程"),
        ("D-18-staff-service-record-flow.png", "附錄｜科辦服務電子獎狀流程"),
    ]
    for filename, title in appendix:
        diagram_slide(prs, filename, title, "原圖取自 docs/images，作為目前規格與實作的對照版本")

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    make_deck()

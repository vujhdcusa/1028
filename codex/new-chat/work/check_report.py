from pathlib import Path
from docx import Document

p = Path(r"C:\Users\lisa2\Documents\Codex\2026-09-14\new-chat\outputs\資管科學習歷程系統_專題成果報告書.docx")
d = Document(str(p))
text = "\n".join(x.text for x in d.paragraphs)
bad = ["**", "####", "more_horiz", "Studio 面板", "graduation", "[1]", "[3]", "[5]"]
print("paragraphs", len(d.paragraphs))
print("tables", len(d.tables))
print("bad_found", {b: text.count(b) for b in bad if b in text})
print("headings", [x.text for x in d.paragraphs if x.style.name.startswith("Heading")][:30])
print("size", p.stat().st_size)

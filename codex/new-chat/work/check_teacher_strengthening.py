from pathlib import Path
from docx import Document

p = Path(r"D:\democase\ai\codex\資管科學習歷程系統_專題成果報告書.docx")
d = Document(str(p))
text = "\n".join(paragraph.text for paragraph in d.paragraphs)
for table in d.tables:
    for row in table.rows:
        text += "\n" + "\t".join(cell.text for cell in row.cells)

checks = [
    "教師輔導省力化",
    "班級學習歷程總覽",
    "風險優先輔導名單",
    "一鍵產生輔導建議",
    "批次催繳通知",
    "面談紀錄與追蹤",
    "班級報表匯出",
    "教師輔導減負模組",
    "Mentor_Followups",
    "Notification_Templates",
]

print("file", p)
print("size", p.stat().st_size)
print("paragraphs", len(d.paragraphs))
print("tables", len(d.tables))
print("missing", [item for item in checks if item not in text])

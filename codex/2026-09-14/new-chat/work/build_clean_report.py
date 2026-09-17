from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = r"C:\Users\lisa2\Documents\Codex\2026-09-14\new-chat\outputs\資管科學習歷程系統_專題成果報告書.docx"


def set_font(run, size=12, bold=False):
    run.font.name = "Microsoft JhengHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft JhengHei")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft JhengHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_para(paragraph, size=12, bold=False, align=None, before=0, after=6, line=1.25):
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    for run in paragraph.runs:
        set_font(run, size, bold)


def add_p(doc, text="", size=12, bold=False, align=None, before=0, after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size, bold)
    set_para(p, size, bold, align, before, after)
    return p


def add_h1(doc, text):
    p = doc.add_heading(level=1)
    p.add_run(text)
    set_para(p, 16, True, before=10, after=8)
    return p


def add_h2(doc, text):
    p = doc.add_heading(level=2)
    p.add_run(text)
    set_para(p, 14, True, before=6, after=6)
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    set_font(r, 10.5, bold)
    set_para(p, 10.5, bold, align, after=2, line=1.15)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        shade_cell(hdr[i], "1F4E79")
        set_cell_text(hdr[i], h, True, WD_ALIGN_PARAGRAPH.CENTER)
        for run in hdr[i].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    set_borders(table)
    add_p(doc, "", after=4)
    return table


doc = Document()
for section in doc.sections:
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

styles = doc.styles
for name in ("Normal", "Heading 1", "Heading 2", "Title"):
    style = styles[name]
    style.font.name = "Microsoft JhengHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    style.font.color.rgb = RGBColor(0, 0, 0)

add_p(doc, "康寧學校財團法人康寧大學", 16, True, WD_ALIGN_PARAGRAPH.CENTER, after=8)
add_p(doc, "資訊管理科", 16, True, WD_ALIGN_PARAGRAPH.CENTER, after=8)
add_p(doc, "畢業專題", 18, True, WD_ALIGN_PARAGRAPH.CENTER, after=24)
add_p(doc, "資管科智慧學習歷程系統", 24, True, WD_ALIGN_PARAGRAPH.CENTER, after=36)
add_p(doc, "專 題 生：111534105 田心于", 13, False, WD_ALIGN_PARAGRAPH.CENTER)
add_p(doc, "　　　　　111534107 張馨羽", 13, False, WD_ALIGN_PARAGRAPH.CENTER)
add_p(doc, "　　　　　111534124 林承諺", 13, False, WD_ALIGN_PARAGRAPH.CENTER)
add_p(doc, "指導老師：李栢松 老師", 13, False, WD_ALIGN_PARAGRAPH.CENTER, after=36)
add_p(doc, "中 華 民 國 116 年 06 月", 13, False, WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

add_h1(doc, "目錄")
toc = [
    "一、緒論",
    "　第一節 開發背景",
    "　第二節 開發動機",
    "　第三節 開發目的",
    "二、開發環境",
    "　第一節 系統開發環境及使用工具軟體說明",
    "三、系統規格書",
    "　第一節 使用者需求",
    "　第二節 使用者案例圖",
    "　第三節 使用者流程圖",
    "四、需求分析",
    "　第一節 功能性需求",
    "　第二節 系統架構圖",
    "五、資料表結構",
    "六、開發時程",
    "七、系統展示與功能說明",
    "八、工作分配",
    "九、參考資料",
]
for item in toc:
    add_p(doc, item, 12, False, after=3)
doc.add_page_break()

add_h1(doc, "一、緒論")
add_h2(doc, "第一節 開發背景")
add_p(doc, "隨著資訊科技發展，教育行政與學術評量逐漸朝向數位化與自動化。五專學生從專一至專五累積的學習歷程、備審資料、程式作品集與實習成果，常散落於個人設備、雲端資料夾或紙本文件中，造成學生整理不易、教師輔導困難，也增加科辦行政彙整與審查的負擔。")
add_p(doc, "在二技推甄、實習媒合或專題初審等情境中，科辦與教師需要人工收件、建立清單、分派評審委員並追蹤審查進度。若缺乏系統化控管，容易出現資料重複、審查進度不透明、分派未避嫌或評分紀錄難以追溯等問題。因此，本專題規劃建置一套資管科智慧學習歷程系統，以集中管理學生資料並輔助教師與行政單位進行審查決策。")
add_h2(doc, "第二節 開發動機")
add_p(doc, "現有學習歷程平台多偏向靜態檔案保存，較難驗證學生作品與履歷主張之間的一致性。例如學生在履歷中描述具備程式開發能力，但審查者仍需另行開啟 GitHub、作品連結或備審檔案進行人工比對。此流程耗時且容易因資料格式不一而影響審查效率。")
add_p(doc, "資管科學生具有長期專業養成與實作成果，若能將學習歷程、GitHub API 資料、OCR 成績單辨識與 AI 一致性審計整合在同一平台，將能協助學生更完整呈現成果，也讓教師與審查委員快速掌握學生能力、學分狀態與推甄準備程度。")
add_h2(doc, "第三節 開發目的")
add_p(doc, "本專題旨在建置一套整合數位歷程管理、智慧避嫌分派、線上評分鎖定、AI 一致性審計、OCR 成績單辨識與二技落點試算的智慧平台。系統第一階段以專題展示與核心流程可運作為目標，延伸功能則作為後續擴充方向。")
for t in [
    "一站式數位歷程管理：學生可建立基本資料、上傳備審 PDF，並綁定 GitHub 與 Canva 等作品連結。",
    "科辦智慧避嫌與行政自動化：科辦可建立審查專案，系統依規則避免重複分派並排除指導老師。",
    "委員線上閱覽與評分鎖定：審查委員可線上預覽資料、填寫評分與評語，送出後鎖定並保留異動紀錄。",
    "AI 輔助一致性檢核：系統可作為延伸功能，比對履歷主張與 GitHub 實作紀錄，產生風險標籤與面試提問建議。",
    "OCR 成績單與二技落點試算：系統可辨識成績單資料，輔助檢查學分缺漏，並試算二技推甄加權後的參考分數。",
]:
    add_p(doc, t, 12, after=3)

add_h1(doc, "二、開發環境")
add_h2(doc, "第一節 系統開發環境及使用工具軟體說明")
add_p(doc, "本系統規劃採三層式網頁架構，前端負責操作介面與資料呈現，後端以 RESTful API 處理業務邏輯，資料庫集中保存帳號、歷程、分派、評分與審計紀錄。")
add_table(doc, ["類別", "工具 / 技術", "用途說明"], [
    ["作業系統", "Windows 11", "系統開發、測試與本機部署環境。"],
    ["網頁編輯軟體", "Visual Studio Code", "原始碼編輯、Git 版本控制與 API 偵錯。"],
    ["前端技術", "HTML5、CSS3、JavaScript", "建立頁面結構、互動操作與非同步資料交換。"],
    ["前端框架", "Bootstrap 5 或 Tailwind CSS", "製作響應式介面與一致化元件樣式。"],
    ["資料視覺化", "Chart.js", "呈現審查進度、成績趨勢、GitHub 語言比例與推甄試算圖表。"],
    ["後端框架", "Python 3.12、Flask", "建置 API、處理權限、分派、評分與資料驗證。"],
    ["資料庫", "MySQL 8.0", "保存系統核心資料，並以關聯式設計維持資料一致性。"],
    ["AI 模組", "Google Gemini API", "規劃用於 OCR 辨識、履歷一致性分析與面試問題建議。"],
    ["第三方 API", "GitHub REST API", "擷取作品語言比例與 commit 紀錄，作為能力佐證參考。"],
    ["資安驗證", "bcrypt / Argon2id、WebAuthn", "密碼雜湊保存，並預留無密碼登入擴充介面。"],
])

add_h1(doc, "三、系統規格書")
add_h2(doc, "第一節 使用者需求")
add_table(doc, ["使用者", "功能需求", "說明"], [
    ["學生", "學習歷程建置", "填寫基本資料、上傳備審 PDF、管理 GitHub 與 Canva 作品連結。"],
    ["學生", "成績單與推甄試算", "上傳成績單供 OCR 辨識，查看畢業學分預警與二技推甄參考試算。"],
    ["學生", "審查進度查詢", "查看未提交、已分派、審查中、已完成等狀態。"],
    ["科辦行政", "帳號與權限管理", "管理學生、導師、審查委員與行政人員帳號。"],
    ["科辦行政", "審查專案管理", "建立審查專案、設定截止日並執行智慧避嫌分派。"],
    ["科辦行政", "評分解鎖與紀錄", "必要時重開已鎖定評分，並保存修改理由與審計日誌。"],
    ["審查委員", "線上評審", "免下載預覽備審 PDF 與作品連結，填寫 0 至 100 分與評語。"],
    ["審查委員", "AI 輔助檢核", "查看一致性風險標籤與面試提問建議，作為審查參考。"],
    ["班導師 / 指導老師", "班級看板", "查看全班缺件、審查進度、學分缺漏與風險預警。"],
    ["班導師 / 指導老師", "催繳提醒", "依學生狀態發送站內或 Email 催繳通知。"],
])
add_h2(doc, "第二節 使用者案例圖")
for line in [
    "學生：登入系統、管理備審檔案、綁定作品連結、上傳成績單、查看審查進度與二技試算結果。",
    "科辦行政：建立審查專案、執行智慧避嫌分派、管理評分解鎖、匯出審查報表。",
    "審查委員：查看分派學生、預覽備審資料、參考 AI 一致性報告、送出評分與評語。",
    "班導師 / 指導老師：查看班級學習歷程進度、掌握缺件與學分預警、協助學生補件。",
]:
    add_p(doc, line, 12, after=4)
add_h2(doc, "第三節 使用者流程圖")
for line in [
    "1. 學生登入系統後，建立個人學習歷程並上傳備審 PDF、成績單與作品連結。",
    "2. 系統進行基本格式檢查，並於延伸功能中執行 OCR 成績單辨識、GitHub API 資料擷取與 AI 一致性分析。",
    "3. 科辦建立審查專案，系統依分派規則檢查重複分派與指導老師避嫌。",
    "4. 審查委員線上閱覽學生資料，填寫分數與評語後送出，系統將評分狀態鎖定。",
    "5. 系統彙整審查結果，產出統計報表、畢業學分預警與二技推甄落點試算結果。",
]:
    add_p(doc, line, 12, after=4)

add_h1(doc, "四、需求分析")
add_h2(doc, "第一節 功能性需求")
modules = [
    ("帳號安全性與身分驗證模組", "支援學生、科辦行政、審查委員與導師等角色。密碼以 bcrypt 或 Argon2id 動態加鹽雜湊保存，並預留 WebAuthn FIDO2 生物辨識登入介面作為延伸功能。"),
    ("學生歷程建置與 OCR 成績單模組", "提供備審 PDF 上傳、作品連結保存與學習成果管理。OCR 成績單辨識可作為延伸功能，將辨識結果與資管科修業科目表比對，顯示畢業學分缺漏預警。"),
    ("科辦智慧避嫌分派模組", "科辦建立專案名稱與截止日後，系統檢查同一專案中學生與委員是否已存在分派紀錄，並避免將學生分派給其指導老師。"),
    ("委員線上評審與鎖定機制模組", "審查委員可線上預覽 PDF 與作品連結，評分欄位限制 0 至 100 分。評分送出後狀態改為 locked；若需修改，必須由科辦重開並記錄原因。"),
    ("AI 一致性審計與面試智囊模組", "系統可在去識別化後串接 AI 模型，參考 GitHub 語言比例與 commit 歷史，比對履歷宣稱與實作證據，輸出一致、部分一致或缺乏佐證等風險標籤，並產生面試提問建議。"),
    ("二技推甄落點智慧算分模組", "系統可建立二技簡章配分資料，依書審、面試與乙級證照加成進行試算。此結果定位為輔助參考，不取代正式招生簡章與人工審查。"),
]
for title, body in modules:
    add_p(doc, title + "：", 12, True, after=2)
    add_p(doc, body, 12, after=6)
add_h2(doc, "第二節 系統架構圖")
add_p(doc, "系統架構分為表現層、應用邏輯層、資料庫層與外部服務層。表現層提供學生、科辦、委員與導師操作介面；應用邏輯層負責登入驗證、檔案管理、分派規則、評分鎖定與試算邏輯；資料庫層保存使用者、歷程、審查專案與紀錄；外部服務層包含 Gemini API 與 GitHub REST API。")
add_table(doc, ["層級", "主要元件", "說明"], [
    ["表現層", "學生 Dashboard、科辦管理台、委員評審頁、導師戰情板", "提供各角色對應操作與資料呈現。"],
    ["應用邏輯層", "Flask REST API、權限驗證、避嫌分派、評分鎖定、算分引擎", "處理核心業務流程與資料驗證。"],
    ["資料庫層", "MySQL 8.0", "保存使用者、歷程、分派、評分、審計與課程規則資料。"],
    ["外部服務層", "Gemini API、GitHub REST API", "支援 OCR、AI 分析與作品資料擷取。"],
    ["檔案儲存層", "uploads 目錄或雲端儲存", "保存備審 PDF、成績單圖片與相關附件。"],
])

add_h1(doc, "五、資料表結構")
add_p(doc, "本系統資料庫採用 MySQL 8.0，核心資料表規劃如下。第一版以可展示與可擴充為原則，保留 AI 與二技試算所需欄位，實作時可依實際功能逐步擴充。")
add_table(doc, ["資料表", "用途", "主要欄位"], [
    ["Users", "使用者資料表", "user_id、username、password_hash、real_name、role、advisor_id、status"],
    ["Portfolios", "學習歷程資料表", "portfolio_id、student_id、pdf_path、github_url、canva_url、claimed_skills、status"],
    ["Curriculum_Rules", "修業科目與學分表", "rule_id、course_name、category、required_credits、replace_course_id"],
    ["Review_Projects", "審查專案表", "project_id、project_name、deadline、created_by、status"],
    ["Project_Assignments", "分派紀錄表", "assignment_id、project_id、student_id、reviewer_id、assigned_at"],
    ["Review_Logs", "評分與審計紀錄表", "log_id、assignment_id、score、comment、is_locked、changed_reason、created_at"],
    ["AI_Audit_Reports", "AI 一致性審計表", "audit_id、portfolio_id、risk_label、confidence、summary、interview_questions"],
])

add_h1(doc, "六、開發時程")
add_table(doc, ["工作項目", "主要內容", "預估期間"], [
    ["主題確認與需求訪談", "確認使用者角色、問題情境與展示重點。", "第 1 至 2 週"],
    ["系統分析與資料庫設計", "設計資料表、權限規則與主要流程。", "第 3 至 5 週"],
    ["前端介面設計", "完成學生、科辦、委員與導師主要頁面。", "第 6 至 9 週"],
    ["後端功能開發", "完成登入、歷程管理、分派、評分與報表 API。", "第 8 至 14 週"],
    ["AI 與外部 API 整合", "開發 OCR、GitHub API 與一致性檢核雛形。", "第 12 至 16 週"],
    ["測試與修正", "進行功能測試、權限測試與展示流程修正。", "第 17 至 19 週"],
    ["文件與簡報整理", "完成成果報告書、簡報與展示腳本。", "第 20 週"],
])

add_h1(doc, "七、系統展示與功能說明")
for line in [
    "登入頁面：使用者輸入帳號密碼後登入系統，系統依角色導向對應首頁。",
    "學生 Dashboard：顯示歷程完成度、檔案上傳狀態、審查進度、學分預警與二技試算摘要。",
    "學習歷程管理：學生可新增、修改與送出備審資料，並管理 GitHub 與 Canva 作品連結。",
    "科辦管理台：可建立審查專案、設定截止日、執行智慧避嫌分派並查看整體進度。",
    "委員評審頁：可線上預覽學生資料、參考 AI 檢核結果並填寫分數與評語。",
    "導師戰情板：呈現班級學生缺件、審查狀態、學分缺漏與風險預警排行。",
    "統計報表：提供審查完成率、平均分數、缺件統計與匯出功能，方便專題展示與行政使用。",
]:
    add_p(doc, line, 12, after=4)

add_h1(doc, "八、工作分配")
add_table(doc, ["學號", "姓名", "工作內容"], [
    ["111534105", "田心于", "系統需求整理、前端頁面設計、報告書撰寫與展示資料整理。"],
    ["111534107", "張馨羽", "資料庫規劃、後端 API 開發、權限與分派流程設計。"],
    ["111534124", "林承諺", "AI 模組規劃、GitHub API 串接、系統測試與簡報製作。"],
])

add_h1(doc, "九、參考資料")
for ref in [
    "Google Gemini API 官方文件",
    "GitHub REST API 官方文件",
    "MySQL 8.0 Reference Manual",
    "Flask 官方文件",
    "Chart.js 官方文件",
    "教育部學習歷程檔案相關說明資料",
]:
    add_p(doc, ref, 12, after=3)

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("資管科智慧學習歷程系統")
    set_font(run, 9)

doc.save(OUT)
print(OUT)

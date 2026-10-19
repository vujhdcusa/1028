"""產出完整備審資料（Word）。章節依目標校系加權排序，內容只用學生已確認的資料。"""

from io import BytesIO

from docx import Document
from docx.shared import Pt

from app.models import CredentialRecord, LearningCourseRecord, PackageReview, Portfolio
from app.services.autobiography_service import CONFIRMED_PORTFOLIO_STATUSES, latest_draft
from app.services.gap_service import analyze_gap, chapter_order

REVIEW_TEXT = {
    "confirmed": "導師已確認",
    "pending": "等待導師確認中",
    "needs_revision": "導師建議修改後再送出",
}


def latest_review(student, program):
    return (PackageReview.query.filter_by(student_id=student.id, rule_version_id=program.id)
            .order_by(PackageReview.requested_at.desc(), PackageReview.id.desc()).first())


def _semester_text(year, semester):
    parts = []
    if year:
        parts.append(f"專{'一二三四五'[year - 1]}" if 1 <= year <= 5 else f"{year} 學年度")
    if semester:
        parts.append({1: "上學期", 2: "下學期"}.get(semester, str(semester)))
    return " ".join(parts)


def _chapter_body(document, student, chapter, program):
    codes = chapter["keys"]
    if chapter["code"] == "02":
        document.add_paragraph("請附上報名資格證明（依當年度簡章規定）。")
        return
    if "grades" in codes:
        courses = LearningCourseRecord.query.filter_by(student_id=student.id, confirmation_status="confirmed").order_by(
            LearningCourseRecord.school_year, LearningCourseRecord.id).all()
        if not courses:
            document.add_paragraph("（尚未提供成績資料）")
            return
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ["年級", "學期", "課程", "學分", "分數"]):
            cell.text = text
        for course in courses:
            row = table.add_row().cells
            row[0].text = str(course.school_year or course.grade_level or "")
            row[1].text = course.semester or ""
            row[2].text = course.course_name
            row[3].text = "" if course.credits is None else str(course.credits)
            row[4].text = str(course.score) if course.score is not None else (course.status_mark or "")
        document.add_paragraph("另附學校核發之歷年成績單正本。")
        return
    if "credential" in codes:
        credentials = CredentialRecord.query.filter_by(student_id=student.id, confirmation_status="confirmed").order_by(CredentialRecord.issue_date).all()
        for credential in credentials:
            detail = "、".join(part for part in (credential.level, credential.issuing_organization,
                                                credential.issue_date.isoformat() if credential.issue_date else None) if part)
            document.add_paragraph(f"{credential.credential_name}（{detail}）" if detail else credential.credential_name,
                                   style="List Bullet")
        if not credentials:
            document.add_paragraph("（尚未提供證照資料）")
        return
    if chapter["code"] == "04":
        # 自傳：有學生儲存過的「自傳小幫手」版本就放進來，沒有就留撰寫提示
        draft = latest_draft(student, program)
        if draft:
            for paragraph in draft.content_json.get("paragraphs", []):
                if paragraph.get("text"):
                    document.add_paragraph(paragraph["text"])
            return
        document.add_paragraph("請在此撰寫。可用系統的「自傳小幫手」產生草稿，或參考「我的五年」中的活動、競賽與幹部經歷。")
        return
    if "autobiography" in codes:
        document.add_paragraph("請在此撰寫讀書計畫：進二技後的學習目標、想修的課程與畢業後規劃。")
        return
    if "project" in codes:
        document.add_paragraph("請附上專題或作品集，說明個人負責的部分與成果。")
        return
    records = Portfolio.query.filter(Portfolio.student_id == student.id, Portfolio.category.in_(codes),
                                     Portfolio.status.in_(CONFIRMED_PORTFOLIO_STATUSES)).order_by(
        Portfolio.school_year, Portfolio.id).all()
    for record in records:
        when = _semester_text(record.school_year, record.semester)
        document.add_paragraph(f"{record.title}{'（' + when + '）' if when else ''}", style="List Bullet")
        if record.description:
            document.add_paragraph(record.description)
    if not records:
        document.add_paragraph("（尚未提供資料）")


def build_package_docx(student, program):
    review = latest_review(student, program)
    gap = analyze_gap(student, program)
    chapters = chapter_order(program)
    profile = student.student_profile

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Microsoft JhengHei"
    style.font.size = Pt(11)

    document.add_heading("備審資料", level=0)
    document.add_paragraph(f"申請：{program.school_name} {program.department_name}（志願代碼 {program.program_code}，{program.admission_year} 學年度）")
    document.add_paragraph(f"學生：{student.display_name}" + (f"　學號：{profile.student_no}　班級：{profile.class_name or ''}" if profile else ""))
    document.add_paragraph(f"導師確認：{REVIEW_TEXT.get(review.status, '尚未送導師確認') if review else '尚未送導師確認'}")
    if gap["readiness"] is not None:
        document.add_paragraph(f"資料準備度：{gap['readiness']}%（依校系加權計算，僅供參考）")

    for position, chapter in enumerate(chapters, start=1):
        # 印出來的章節編號依實際順序 01、02、03…，不保留原始編號
        document.add_heading(f"{position:02d}　{chapter['title']}", level=1)
        if chapter["code"] == "01":
            document.add_paragraph("目錄（依本校系加權排序，加權高的項目排在前面）：")
            for position, item in enumerate(chapters[1:], start=1):
                note = f"（相關加權 {item['weight']}%）" if item["weight"] else ""
                document.add_paragraph(f"第{position}部分：{item['title']}{note}")
            continue
        _chapter_body(document, student, chapter, program)

    document.add_paragraph("本文件內容皆為學生確認之資料；正式規定以當年度招生簡章與各校系分則為準。")
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream

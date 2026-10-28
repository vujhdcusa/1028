"""自傳「兩邊都存」的學校端備份：每建立一版自傳，就另存一份 Word 檔到學校的檔案存放（使用者 10-28 決定）。

- 存放位置交給 storage_service：管理員連上 Google 雲端硬碟時存到「智慧學習歷程系統上傳檔」資料夾，
  否則存本機 UPLOAD_FOLDER/autobiography/<學生 id>/。
- 位置寫進該版 content_json["backup_location"]；備份失敗只記 log、位置留空，不擋生成或存檔。
- 正式版本仍在資料庫；這份 Word 只是備份。學生自己的 Google 文件另走 google_docs_service。
"""
import io
import re
from datetime import datetime, timedelta, timezone

from docx import Document
from flask import current_app

from app.services import storage_service

FOLDER = "autobiography"
TAIPEI = timezone(timedelta(hours=8))

# 狀態代碼 → 給使用者看的文字。不顯示本機路徑或 Drive 內部位置。
STATUS_LABELS = {
    "cloud": f"已備份到學校雲端（{storage_service.FOLDER_NAME}）",
    "local": "已備份到本機（學校雲端未連線）",
    "failed": "備份失敗",
}

_UNSAFE = re.compile(r"[/\\:*?\"<>|\s]+")


def _safe(text):
    return _UNSAFE.sub("", str(text or "")) or "未命名"


def _account(student):
    profile = getattr(student, "student_profile", None)
    return (profile.student_no if profile and profile.student_no else None) or student.username


def filename_for(student, program, version):
    return f"自傳_{_safe(_account(student))}_{_safe(program.school_name + program.department_name)}_第{version}版.docx"


def build_docx(student, program, content):
    """標題含學生姓名、目標校系、第幾版、時間；每段一個段落（段落標題另起小標）。"""
    document = Document()
    now = datetime.now(TAIPEI).strftime("%Y-%m-%d %H:%M")
    document.add_heading(f"自傳｜{student.display_name}｜{program.school_name} {program.department_name}", level=1)
    document.add_paragraph(f"第 {content.get('version')} 版｜{now}（台北時間）")
    for paragraph in content.get("paragraphs", []):
        if paragraph.get("title"):
            document.add_heading(paragraph["title"], level=2)
        document.add_paragraph(paragraph.get("text", ""))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def backup(student, program, content):
    """存一份 Word 並回傳位置；任何錯誤都只記 log，回傳空字串。"""
    try:
        data = build_docx(student, program, content)
        return storage_service.save(data, student.id, FOLDER, filename_for(student, program, content.get("version")))
    except Exception as error:  # 備份失敗不能讓生成或存檔失敗
        current_app.logger.warning("自傳備份失敗（學生 %s 第 %s 版）：%s", student.id, content.get("version"), error)
        return ""


def status_of(output):
    """這一版的備份狀態：cloud／local／failed；舊資料沒有備份欄位時回傳 None。"""
    content = output.content_json or {}
    if "backup_location" not in content:
        return None
    location = content.get("backup_location") or ""
    if not location:
        return "failed"
    return "cloud" if location.startswith(storage_service.DRIVE_PREFIX) else "local"


def status_label(output):
    return STATUS_LABELS.get(status_of(output))

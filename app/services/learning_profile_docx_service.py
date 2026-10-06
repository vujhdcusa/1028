from io import BytesIO

from docx import Document


OUTPUT_TITLES = {
    "profile_summary": "學習歷程摘要",
    "curriculum_gap": "課程規則比對與待確認項目",
    "autobiography": "自傳草稿",
    "study_plan": "學習與升學規劃草稿",
}


def output_to_docx_bytes(output, student_name):
    document = Document()
    document.add_heading(OUTPUT_TITLES.get(output.output_type, "學習歷程輸出"), level=1)
    document.add_paragraph(f"學生：{student_name}")
    content = output.content_json
    for key, value in content.items():
        if key == "ability_domains" and isinstance(value, list):
            document.add_heading("能力領域", level=2)
            for item in value:
                document.add_paragraph(
                    f"{item['name']}：平均 {item['average_score']} 分（{item['course_count']} 門課）",
                    style="List Bullet",
                )
        elif isinstance(value, list):
            document.add_heading(key, level=2)
            for item in value:
                document.add_paragraph(str(item), style="List Bullet")
        else:
            document.add_paragraph(f"{key}：{value}")
    document.add_paragraph("資料來源與規則版本已隨本次輸出保留於系統。")
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream

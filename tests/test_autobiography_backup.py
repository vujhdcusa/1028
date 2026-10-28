"""自傳「兩邊都存」：每一版自傳另存一份 Word 到學校端（Drive 或本機），失敗不擋生成或存檔。"""
from datetime import date
import io

import pytest
from docx import Document

from app import create_app
from app.extensions import db
from app.models import AdmissionProgram, LearningProfileOutput, StudentProfile, User
from app.services import autobiography_service, storage_service
from app.services.admission_service import import_seed_programs
from app.services.journey_service import academic_year_of
from tests.test_storage_service import FakeDriveClient

ANSWERS = {"motivation": "專二做網頁設計時發現自己喜歡把需求變成系統。", "proud": "黑客松時負責後端 API。",
           "challenge": "", "future": "想學資料分析。"}


def make_app(tmp_path, **config):
    app = create_app({"SQLALCHEMY_DATABASE_URI": "sqlite:///:memory:", "TESTING": True,
                      "UPLOAD_FOLDER": str(tmp_path / "uploads"), **config})
    with app.app_context():
        db.create_all()
        import_seed_programs()
        for username in ("stu", "other"):
            user = User(username=username, display_name=f"學生{username}", role="student")
            user.set_password("pw")
            db.session.add(user)
        db.session.flush()
        for username, no in (("stu", "111534100"), ("other", "111534199")):
            user = User.query.filter_by(username=username).one()
            db.session.add(StudentProfile(user_id=user.id, student_no=no, class_name="資管科五甲",
                                          admission_year=academic_year_of(date.today()) - 4))  # 專五
        db.session.commit()
    return app


@pytest.fixture()
def app(tmp_path):
    return make_app(tmp_path)


def program_id(app):
    with app.app_context():
        return AdmissionProgram.query.filter_by(program_code="208003").one().id


def login(app, username="stu"):
    client = app.test_client()
    client.post("/login", data={"username": username, "password": "pw"})
    client.post("/student/targets", data={"program_id": program_id(app)})
    return client


def generate(app, client):
    response = client.post("/student/autobiography/generate", data={"program_id": program_id(app), **ANSWERS})
    assert response.status_code == 302
    with app.app_context():
        return LearningProfileOutput.query.filter_by(output_type="autobiography_ai").order_by(
            LearningProfileOutput.id.desc()).first().id


def content_of(app, output_id):
    with app.app_context():
        return dict(db.session.get(LearningProfileOutput, output_id).content_json)


def test_generate_backs_up_docx_locally(app):
    client = login(app)
    output_id = generate(app, client)

    content = content_of(app, output_id)
    location = content["backup_location"]
    assert location and not location.startswith("drive:")
    with app.app_context():
        data = storage_service.read_bytes(location)
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "學生stu" in text
    assert "第 1 版" in text
    assert "專二做網頁設計" in text
    assert "自傳_111534100_" in location and "第1版.docx" in location
    assert "autobiography" in location

    page = client.get(f"/student/autobiography?program_id={program_id(app)}").get_data(as_text=True)
    assert "已備份到本機（學校雲端未連線）" in page
    assert location not in page
    assert str(app.config["UPLOAD_FOLDER"]) not in page


def test_backup_failure_does_not_block_generation(app, monkeypatch):
    def boom(*args, **kwargs):
        raise storage_service.StorageUnavailable("壞掉了")

    monkeypatch.setattr(storage_service, "save", boom)
    client = login(app)
    output_id = generate(app, client)

    content = content_of(app, output_id)
    assert content["paragraphs"]
    assert not content["backup_location"]
    page = client.get(f"/student/autobiography?program_id={program_id(app)}").get_data(as_text=True)
    assert "備份失敗" in page
    versions = client.get(f"/student/autobiography/versions?program_id={program_id(app)}").get_data(as_text=True)
    assert "備份失敗" in versions


def _count_saves(monkeypatch):
    calls = []
    real = storage_service.save

    def counting(data, owner_id, folder, filename):
        calls.append((owner_id, folder, filename))
        return real(data, owner_id, folder, filename)

    monkeypatch.setattr(storage_service, "save", counting)
    return calls


def test_saved_edits_are_backed_up(app, monkeypatch):
    calls = _count_saves(monkeypatch)
    client = login(app)
    output_id = generate(app, client)
    count = len(content_of(app, output_id)["paragraphs"])

    client.post(f"/student/autobiography/{output_id}/save",
                data={f"paragraph_{i}": f"改過第{i + 1}段" for i in range(count)})

    assert len(calls) == 2
    assert all(folder == "autobiography" for _, folder, _ in calls)
    assert calls[1][2].endswith("第2版.docx")
    with app.app_context():
        newest = LearningProfileOutput.query.order_by(LearningProfileOutput.id.desc()).first()
        assert newest.content_json["backup_location"]
        text = "\n".join(p.text for p in Document(io.BytesIO(
            storage_service.read_bytes(newest.content_json["backup_location"]))).paragraphs)
        assert "改過第1段" in text


def test_imported_version_is_backed_up(app, monkeypatch):
    calls = _count_saves(monkeypatch)
    client = login(app)
    output_id = generate(app, client)
    with app.app_context():
        output = db.session.get(LearningProfileOutput, output_id)
        student = User.query.filter_by(username="stu").one()
        new_version = autobiography_service.save_imported(
            output, [{"title": "學習動機", "text": "從 Google 取回的文字", "sources": []}], student, origin="google_docs")
        assert new_version.content_json["backup_location"]
    assert len(calls) == 2


def test_drive_mode_backs_up_to_school_drive(tmp_path):
    fake = FakeDriveClient()
    app = make_app(tmp_path, STORAGE_BACKEND="google_drive", GOOGLE_OAUTH_CLIENT_ID="cid",
                   GOOGLE_OAUTH_CLIENT_SECRET="secret",
                   GOOGLE_STORAGE_REDIRECT_URI="http://localhost/admin/storage/google/callback",
                   GOOGLE_STORAGE_CLIENT=fake)
    storage_service.token_holder(app).set({"access_token": "tok-1", "expires_in": 3600, "refresh_token": "r"})
    client = login(app)
    output_id = generate(app, client)

    location = content_of(app, output_id)["backup_location"]
    assert location.startswith("drive:")
    (stored,) = fake.files.values()
    assert stored["app_properties"]["slp_folder"] == "autobiography"
    assert "自傳_111534100_" in stored["name"]
    assert stored["mime"].endswith("wordprocessingml.document")
    page = client.get(f"/student/autobiography?program_id={program_id(app)}").get_data(as_text=True)
    assert "已備份到學校雲端" in page


def test_student_sees_only_own_backup_status(app, monkeypatch):
    client = login(app)
    generate(app, client)
    other = login(app, "other")

    def boom(*args, **kwargs):
        raise storage_service.StorageUnavailable("壞掉了")

    monkeypatch.setattr(storage_service, "save", boom)
    generate(app, other)

    mine = client.get(f"/student/autobiography/versions?program_id={program_id(app)}").get_data(as_text=True)
    theirs = other.get(f"/student/autobiography/versions?program_id={program_id(app)}").get_data(as_text=True)
    assert "已備份到本機" in mine and "備份失敗" not in mine
    assert "備份失敗" in theirs and "已備份到本機" not in theirs


def test_page_points_to_own_google_doc_even_when_unconfigured(app):
    client = login(app)
    generate(app, client)

    page = client.get(f"/student/autobiography?program_id={program_id(app)}").get_data(as_text=True)

    assert "在自己的 Google 文件中編輯" in page
    assert "學校尚未設定 Google 連線" in page


def test_admin_storage_page_mentions_autobiography_backup(app):
    with app.app_context():
        admin = User(username="adm", display_name="管理員", role="admin")
        admin.set_password("pw")
        db.session.add(admin)
        db.session.commit()
    client = app.test_client()
    client.post("/login", data={"username": "adm", "password": "pw"})

    page = client.get("/admin/storage").get_data(as_text=True)

    assert "自傳備份" in page
    assert storage_service.FOLDER_NAME in page
